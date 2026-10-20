"""
batsim_to_csv.py - combine Regatron BatSim .batsimdat channel files into one
table per session.

Each .batsimdat file holds one channel as "ticks;value" lines, where ticks are
.NET DateTime ticks (100 ns units since 1 Jan 0001). Files are named
    <SessionName>_<channel>_<part>.batsimdat
e.g. Session1_300920_60SOC_0_1.batsimdat, with channel index:
    0 = Voltage Act. (V)   1 = Current Act. (A)
    2 = State of charge (%)   3 = Internal resistance (mOhm)
(from the session index file). Multi-part channels (_1, _2, ...) are joined.

Columns: time_s (seconds from the first reading of each session), voltage_V,
current_A, internal_resistance_mOhm, state_of_charge_pct.

Output (in an "output" subfolder):
    <SessionName>.csv          one CSV per session
    all_sessions.xlsx          one sheet per session (needs openpyxl)

Incremental mode (default): sessions that already have a sheet in
all_sessions.xlsx are skipped, and only NEW sessions are converted and added
as new sheets (in the order they were recorded). Existing sheets are never
changed. Keep all days' .batsimdat files in the SAME folder.

Usage:
    python batsim_to_csv.py                      # this folder, add new sessions
    python batsim_to_csv.py "C:\\path\\to\\data"   # another folder
    python batsim_to_csv.py --rebuild            # redo everything from scratch

Requires: pandas (pip install pandas openpyxl)

AI-USE ACKNOWLEDGEMENT (MEC4128Z): drafted with Claude (Anthropic),
30 Sep 2026, incremental mode added 1 Oct 2026 - see attached chat log.
Data conversion only; no analysis.
"""

import re
import sys
from pathlib import Path

import pandas as pd

CHANNELS = {
    0: "voltage_V",
    1: "current_A",
    2: "state_of_charge_pct",
    3: "internal_resistance_mOhm",
}
COLUMN_ORDER = ["time_s", "voltage_V", "current_A",
                "internal_resistance_mOhm", "state_of_charge_pct"]
NAME_RE = re.compile(r"^(?P<session>.+)_(?P<chan>\d+)_(?P<part>\d+)\.batsimdat$",
                     re.IGNORECASE)


def read_channel(files):
    """Read one channel (possibly several parts) into a Series indexed by ticks."""
    ticks, values = [], []
    for f in files:
        with open(f, encoding="utf-8", errors="replace") as fh:
            for line in fh:
                line = line.strip()
                if not line or ";" not in line:
                    continue
                t, v = line.split(";", 1)
                try:
                    ticks.append(int(t))          # int keeps full 18-digit precision
                    values.append(float(v))
                except ValueError:
                    continue                       # skip any header/garbage lines
    s = pd.Series(values, index=ticks, dtype="float64")
    return s[~s.index.duplicated(keep="first")].sort_index()


def existing_sheets(xlsx_path: Path):
    """Sheet names already in the workbook (empty set if it does not exist)."""
    if not xlsx_path.exists():
        return set()
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, read_only=True)
    names = set(wb.sheetnames)
    wb.close()
    return names


def convert_folder(folder: Path, rebuild: bool = False):
    groups = {}   # session -> channel -> [(part, path)]
    for f in folder.glob("*.batsimdat"):
        m = NAME_RE.match(f.name)
        if not m:
            print(f"  skipped (name not recognised): {f.name}")
            continue
        chan = int(m["chan"])
        groups.setdefault(m["session"], {}).setdefault(chan, []).append((int(m["part"]), f))

    if not groups:
        print(f"No .batsimdat files found in {folder}")
        return

    out_dir = folder / "output"
    out_dir.mkdir(exist_ok=True)
    xlsx_path = out_dir / "all_sessions.xlsx"
    if rebuild and xlsx_path.exists():
        xlsx_path.unlink()
        print("  --rebuild: old workbook deleted, converting everything")
    try:
        done = existing_sheets(xlsx_path)
    except ImportError:
        done = set()
    tables = {}

    for session in sorted(groups):
        if session[:31] in done:
            print(f"  {session}: already in workbook - skipped")
            continue
        cols = {}
        for chan, parts in groups[session].items():
            name = CHANNELS.get(chan, f"channel_{chan}")
            cols[name] = read_channel([p for _, p in sorted(parts)])
        df = pd.DataFrame(cols)                          # aligned on ticks
        ticks = df.index.to_numpy(dtype="int64")
        df.insert(0, "time_s", (ticks - ticks[0]) / 1e7)
        df = df[[c for c in COLUMN_ORDER if c in df.columns] +
                [c for c in df.columns if c not in COLUMN_ORDER]]
        df.reset_index(drop=True, inplace=True)

        missing = [CHANNELS[c] for c in CHANNELS if CHANNELS[c] not in df.columns]
        n_gaps = int(df.drop(columns=["time_s"]).isna().any(axis=1).sum())
        csv_path = out_dir / f"{session}.csv"
        df.to_csv(csv_path, index=False)
        tables[session] = (ticks[0], df)
        print(f"  {session}: {len(df)} rows, {df['time_s'].iloc[-1]:.1f} s"
              + (f", MISSING {missing}" if missing else "")
              + (f", {n_gaps} rows with a blank channel" if n_gaps else "")
              + f" -> {csv_path.name}")

    if not tables:
        print("  Nothing new to add.")
        return
    # add new sheets in the order the sessions were recorded (first timestamp)
    new = sorted(tables.items(), key=lambda kv: kv[1][0])
    try:
        mode = "a" if xlsx_path.exists() else "w"
        with pd.ExcelWriter(xlsx_path, engine="openpyxl", mode=mode) as xw:
            for session, (_, df) in new:
                df.to_excel(xw, sheet_name=session[:31], index=False)  # Excel 31-char limit
        print(f"  Added {len(new)} sheet(s) to {xlsx_path.name}")
    except ImportError:
        print("  (openpyxl not installed - CSVs written, Excel workbook skipped)")
    except PermissionError:
        print(f"  Could not write {xlsx_path.name} - close it in Excel and run again.")


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--rebuild"]
    rebuild = "--rebuild" in sys.argv[1:]
    folder = Path(args[0]) if args else Path(__file__).parent
    print(f"Converting .batsimdat files in {folder}")
    convert_folder(folder, rebuild)